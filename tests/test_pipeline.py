"""End-to-end pipeline tests against the in-process fake internet (no real network)."""

from __future__ import annotations

import csv
import json
import os
import sqlite3
import xml.etree.ElementTree as ET

import httpx

from webscraper.job.models import parse_job
from webscraper.job.runner import JobControl, JobRunner
from webscraper.job.state import JobState
from webscraper.storage.database import Database

from .conftest import MockWeb


def business_job(**overrides):
    job = {
        "name": "Cape Town software companies",
        "entity": "business",
        "targeting": {
            "required": {"location": {"country": "South Africa", "province": "Western Cape", "city": "Cape Town"},
                         "industry": ["Software"]},
            "preferred": {"keywords": ["SaaS", "cloud"]},
            "excluded": {"business_types": ["Government"]},
        },
        "fields": [
            {"name": "company_name", "required": True},
            {"name": "website", "type": "url", "required": True},
            {"name": "email", "type": "email"},
            {"name": "phone", "type": "phone"},
            {"name": "address", "type": "address"},
            {"name": "city"},
            {"name": "country"},
            {"name": "industry"},
            {"name": "social_links", "type": "array"},
        ],
        "sources": {"mode": "auto", "max_queries": 1},
        "records": {"minimum": 2, "maximum": 10},
        "crawl": {"delay": 0, "concurrency": 4},
        "output": {"format": ["csv", "xlsx", "json", "jsonl", "xml", "sql"], "filename": "cpt",
                   "include_provenance": True},
    }
    job.update(overrides)
    return job


def run(job_dict, web: MockWeb, tmp_path, db=None, control=None):
    job = parse_job(job_dict)
    runner = JobRunner(job, db=db, control=control, output_dir=str(tmp_path / "out"), transport=web.transport)
    return job, runner.run_sync()


def by_name(records, prefix):
    for r in records:
        if str(r["data"].get("company_name", "")).startswith(prefix):
            return r
    raise AssertionError("no record named %s in %s" % (prefix, [r["data"].get("company_name") for r in records]))


def test_site_mode_discovery_directory_enrichment_and_dedup(business_web, tmp_path, no_sleep):
    db = Database(str(tmp_path / "db.sqlite"))
    job, outcome = run(business_job(), business_web, tmp_path, db=db)

    assert outcome.state == JobState.COMPLETED, outcome.error
    valid = [r for r in outcome.records if r["valid"]]
    names = sorted(r["data"]["company_name"] for r in valid)
    # Gamma is excluded ("Government"), DeltaDev has no website, Secret is blocked by robots.txt.
    assert len(valid) == 2, names
    assert any(n.startswith("AlphaSoft") for n in names) and "BetaCloud" in names

    alpha = by_name(valid, "AlphaSoft")
    data = alpha["data"]
    assert data["website"] == "https://www.alphasoft.co.za"
    assert data["email"] == "info@alphasoft.co.za"
    assert data["phone"] == "+27215550101"
    assert data["city"] == "Cape Town"
    assert data["country"] == "South Africa"
    assert data["industry"] == "Software"
    assert "https://www.linkedin.com/company/alphasoft" in data["social_links"]
    assert not any("sharer" in s for s in data["social_links"])
    # Directory listing + company website were merged into one record with both sources.
    assert any("capedirectory.test" in s for s in alpha["sources"])
    assert any("alphasoft.co.za" in s for s in alpha["sources"])
    assert alpha["provenance"]["email"]["method"] == "mailto"

    beta = by_name(valid, "BetaCloud")["data"]
    assert beta["email"] == "hello@betacloud.co.za"  # Cloudflare-protected address decoded
    assert beta["phone"] == "+27215550202"
    assert beta["website"] == "https://betacloud.co.za"  # upgraded: the site redirects to https

    fetched = business_web.fetched()
    assert not any("/private/" in u for u in fetched), "robots.txt must be respected"
    assert not any("facebook.com" in u for u in fetched)
    assert not any("hello@capedirectory" in json.dumps(r["data"]) for r in valid), "directory footer leaked"

    m = outcome.metrics
    assert m.queries_run >= 1 and m.sources_discovered >= 2
    assert m.duplicates_removed >= 1
    assert m.final_unique == 2

    # Outputs in every requested format.
    paths = {os.path.splitext(p)[1]: p for p in outcome.outputs}
    assert set(paths) == {".csv", ".xlsx", ".json", ".jsonl", ".xml", ".sql"}
    with open(paths[".csv"], encoding="utf-8-sig") as fh:
        rows = list(csv.DictReader(fh))
    assert len(rows) == 2 and "Company Name" in rows[0] and "Quality Score" in rows[0]
    doc = json.load(open(paths[".json"], encoding="utf-8"))
    assert doc["count"] == 2 and doc["records"][0]["provenance"]
    assert len(open(paths[".jsonl"], encoding="utf-8").read().strip().splitlines()) == 2
    assert len(ET.parse(paths[".xml"]).getroot().findall("record")) == 2
    conn = sqlite3.connect(":memory:")
    conn.executescript(open(paths[".sql"], encoding="utf-8").read())
    assert conn.execute('SELECT COUNT(*) FROM "cpt"').fetchone()[0] == 2
    from openpyxl import load_workbook

    wb = load_workbook(paths[".xlsx"])
    assert wb.sheetnames == ["Records", "Provenance", "Summary"]
    assert wb["Records"].max_row == 3

    # Everything is persisted for the results view.
    stored = db.records(job.id)
    assert len(stored) == 2
    assert db.get_job(job.id)["state"] == "COMPLETED"
    assert db.provenance(job.id)
    assert db.raw_page_count(job.id) == m.pages_fetched


def test_rejected_records_are_reported_with_reasons(business_web, tmp_path, no_sleep):
    job = business_job()
    job["output"] = {"format": "json", "filename": "with-rejected", "include_rejected": True}
    _, outcome = run(job, business_web, tmp_path)
    rejected = {r["data"].get("company_name"): r["reasons"] for r in outcome.records if not r["valid"]}
    assert any("Gamma" in (n or "") for n in rejected)
    gamma = next(v for k, v in rejected.items() if k and "Gamma" in k)
    assert any("Government" in reason for reason in gamma)
    delta = next(v for k, v in rejected.items() if k and "DeltaDev" in k)
    assert any("website" in reason for reason in delta)


def test_stop_at_minimum_stops_early(business_web, tmp_path, no_sleep):
    job = business_job(records={"minimum": 1, "maximum": 10, "stop_at_minimum": True})
    _, outcome = run(job, business_web, tmp_path)
    assert outcome.state == JobState.COMPLETED
    assert sum(1 for r in outcome.records if r["valid"]) >= 1
    assert outcome.metrics.pages_fetched < 14


def test_partially_completed_when_not_enough_records(business_web, tmp_path, no_sleep):
    job = business_job(records={"minimum": 5, "maximum": 10})
    _, outcome = run(job, business_web, tmp_path)
    assert outcome.state == JobState.PARTIALLY_COMPLETED
    assert outcome.outputs  # still exports what was found


def test_cancel_before_start_exports_nothing_and_is_cancelled(business_web, tmp_path, no_sleep):
    control = JobControl()
    control.cancel()
    _, outcome = run(business_job(), business_web, tmp_path, control=control)
    assert outcome.state == JobState.CANCELLED


def test_search_blocked_is_reported(tmp_path, no_sleep):
    web = MockWeb()
    web.handler("html.duckduckgo.com", lambda r: httpx.Response(202, text="<html>anomaly</html>"))
    db = Database(str(tmp_path / "db.sqlite"))
    job, outcome = run(business_job(), web, tmp_path, db=db)
    assert outcome.state == JobState.PARTIALLY_COMPLETED
    messages = " ".join(e["message"] for e in db.events(job.id))
    assert "rate limiting" in messages and "No sources found" in messages


# ------------------------------------------------------------------------------------------------
# Page mode: products, pagination, JSON-LD, sitemaps, feeds and APIs
# ------------------------------------------------------------------------------------------------

def shop_web(n_products=12, per_page=5) -> MockWeb:
    web = MockWeb()
    pages = (n_products + per_page - 1) // per_page
    for p in range(1, pages + 1):
        items = []
        for i in range((p - 1) * per_page + 1, min(n_products, p * per_page) + 1):
            items.append('<article class="product"><h3><a href="/product/widget-%d">Widget %d</a></h3>'
                         '<span class="price">$%d.99</span></article>' % (i, i, i))
        nxt = '<a href="/shop?page=%d">Next</a>' % (p + 1) if p < pages else ""
        url = "https://shop.test/shop" + ("" if p == 1 else "?page=%d" % p)
        web.html(url, "<html><head><title>Shop</title></head><body><main>%s%s</main>"
                      "<footer><a href='/privacy'>Privacy</a></footer></body></html>" % ("".join(items), nxt))
    for i in range(1, n_products + 1):
        ld = {"@context": "https://schema.org", "@type": "Product", "name": "Widget %d" % i, "sku": "W-%03d" % i,
              "brand": {"@type": "Brand", "name": "Acme"},
              "offers": {"@type": "Offer", "price": "%d.99" % i, "priceCurrency": "USD",
                         "availability": "https://schema.org/InStock"}}
        web.html("https://shop.test/product/widget-%d" % i,
                 '<html><head><title>Widget %d | Shop</title><script type="application/ld+json">%s</script>'
                 '</head><body><h1>Widget %d</h1><table><tr><th>Weight</th><td>%d kg</td></tr></table>'
                 '</body></html>' % (i, json.dumps(ld), i, i))
    return web


def product_job(**overrides):
    job = {
        "entity": "product",
        "sources": {"mode": "manual", "urls": [{"url": "https://shop.test/shop", "type": "listing"}]},
        "fields": [
            {"name": "name", "required": True},
            {"name": "price", "type": "decimal", "required": True},
            {"name": "currency"},
            {"name": "availability"},
            {"name": "brand"},
            {"name": "sku"},
            {"name": "weight", "type": "number"},
            {"name": "source_url", "type": "url"},
        ],
        "records": {"minimum": 12, "maximum": 50},
        "crawl": {"delay": 0},
        "output": {"format": "json", "filename": "shop"},
    }
    job.update(overrides)
    return job


def test_page_mode_listing_pagination_and_detail_pages(tmp_path):
    web = shop_web()
    _, outcome = run(product_job(), web, tmp_path)
    assert outcome.state == JobState.COMPLETED, outcome.error
    valid = [r for r in outcome.records if r["valid"]]
    assert len(valid) == 12
    w3 = next(r["data"] for r in valid if r["data"]["name"] == "Widget 3")
    assert w3 == {"name": "Widget 3", "price": 3.99, "currency": "USD", "availability": "InStock", "brand": "Acme",
                  "sku": "W-003", "weight": 3.0, "source_url": "https://shop.test/product/widget-3"}
    assert not any("privacy" in u for u in web.fetched())
    assert not any(r["data"]["name"] == "Shop" for r in outcome.records)


def test_page_mode_record_selector_on_listing(tmp_path):
    web = shop_web()
    job = product_job()
    job["sources"]["urls"][0]["record_selector"] = "article.product"
    job["fields"] = [{"name": "name", "selector": "h3", "required": True},
                     {"name": "price", "type": "decimal", "selector": ".price"},
                     {"name": "source_url", "type": "url"}]
    job["crawl"]["max_depth"] = 0  # listing pages only (pagination keeps the same depth)
    _, outcome = run(job, web, tmp_path)
    valid = [r for r in outcome.records if r["valid"]]
    assert len(valid) == 12
    assert {r["data"]["source_url"] for r in valid} == {"https://shop.test/product/widget-%d" % i for i in range(1, 13)}
    assert not any("/product/" in u for u in web.fetched())


def test_sitemap_source(tmp_path):
    web = shop_web()
    sitemap = ('<?xml version="1.0"?><urlset xmlns="http://www.sitemaps.org/schemas/sitemap/0.9">%s</urlset>' %
               "".join("<url><loc>https://shop.test/product/widget-%d</loc></url>" % i for i in range(1, 13)))
    web.add("https://shop.test/sitemap.xml", sitemap, content_type="application/xml")
    job = product_job(sources={"mode": "manual", "sitemaps": ["https://shop.test/sitemap.xml"]})
    _, outcome = run(job, web, tmp_path)
    assert sum(1 for r in outcome.records if r["valid"]) == 12
    assert not any(u.startswith("https://shop.test/shop") for u in web.fetched())


def test_rss_feed_source(tmp_path):
    web = shop_web()
    feed = "<rss><channel>%s</channel></rss>" % "".join(
        "<item><title>W%d</title><link>https://shop.test/product/widget-%d</link></item>" % (i, i) for i in range(1, 6))
    web.add("https://shop.test/feed.xml", feed, content_type="application/rss+xml")
    job = product_job(sources={"mode": "manual", "urls": [{"url": "https://shop.test/feed.xml", "type": "feed"}]},
                      records={"minimum": 5, "maximum": 5})
    _, outcome = run(job, web, tmp_path)
    assert sum(1 for r in outcome.records if r["valid"]) == 5


def test_json_api_source_with_pagination(tmp_path):
    web = MockWeb()

    def api(request: httpx.Request) -> httpx.Response:
        page = int(request.url.params.get("page", "1"))
        if page > 2:
            return httpx.Response(200, json={"data": {"items": []}})
        items = [{"id": "%d-%d" % (page, i), "title": "Job %d.%d" % (page, i), "company": {"name": "Acme"},
                  "posted": "2026-0%d-1%d" % (page, i)} for i in range(3)]
        return httpx.Response(200, json={"data": {"items": items}})

    web.handler("api.jobs.test", api)
    job = {
        "entity": "job",
        "sources": {"mode": "manual", "urls": [{"url": "https://api.jobs.test/v1/jobs", "type": "api",
                                                "records_path": "data.items", "page_param": "page",
                                                "field_map": {"job_title": "title", "company": "company.name",
                                                              "date_posted": "posted"}}]},
        "fields": [{"name": "job_title", "required": True}, {"name": "company"},
                   {"name": "date_posted", "type": "date"}],
        "records": {"minimum": 6, "maximum": 6},
        "crawl": {"delay": 0},
        "output": {"format": "csv", "filename": "jobs"},
    }
    _, outcome = run(job, web, tmp_path)
    valid = [r["data"] for r in outcome.records if r["valid"]]
    assert len(valid) == 6
    assert {"job_title": "Job 2.1", "company": "Acme", "date_posted": "2026-02-11"} in valid


def test_blocked_source_is_skipped(tmp_path):
    web = MockWeb()
    web.handler("blocked.test", lambda r: httpx.Response(403, text="Forbidden"))
    job = product_job(sources={"mode": "manual", "urls": [{"url": "https://blocked.test/shop", "type": "listing"}]})
    _, outcome = run(job, web, tmp_path)
    assert outcome.state in (JobState.PARTIALLY_COMPLETED, JobState.SOURCE_BLOCKED)
    assert outcome.metrics.pages_failed >= 1


def test_page_budget_is_respected(tmp_path):
    web = shop_web(n_products=40, per_page=10)
    job = product_job(records={"minimum": 40, "maximum": 40}, crawl={"delay": 0, "max_pages": 8})
    _, outcome = run(job, web, tmp_path)
    assert outcome.metrics.pages_fetched <= 8 + 6  # in-flight requests may finish after the budget is hit
    assert outcome.state == JobState.PARTIALLY_COMPLETED
